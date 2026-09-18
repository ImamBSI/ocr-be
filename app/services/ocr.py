import logging
import re
from pathlib import Path
from typing import Dict, List

import PyPDF2
import pytesseract
from docx import Document
from pdf2image import convert_from_path
from PIL import Image

from app.core.config import settings

logger = logging.getLogger(__name__)


class OCRService:
    """Service untuk OCR processing."""

    @staticmethod
    def extract_text_from_pdf(file_path: str) -> str:
        """
        Extract text dari PDF file.

        Args:
            file_path: Path ke PDF file

        Returns:
            Extracted text
        """
        try:
            logger.info(f"Extracting text from PDF: {file_path}")

            text = ""

            # Coba PyPDF2 terlebih dahulu (faster untuk text-based PDF)
            try:
                with open(file_path, "rb") as file:
                    pdf_reader = PyPDF2.PdfReader(file)
                    for page in pdf_reader.pages:
                        text += page.extract_text() + "\n"

                if text.strip():
                    return text
            except Exception as e:
                logger.warning(
                    f"PyPDF2 extraction failed: {e}, trying OCR..."
                )

            # Fallback ke OCR jika PDF text-based extraction gagal (scanned PDF)
            images = convert_from_path(
                file_path,
                dpi=settings.PDF_DPI,
                first_page=1,
                last_page=None,
            )

            for i, image in enumerate(images):
                logger.info(f"Processing PDF page {i + 1}")
                page_text = pytesseract.image_to_string(image)
                text += page_text + "\n"

            return text.strip()

        except Exception as e:
            logger.error(f"Error extracting text from PDF: {str(e)}")
            raise

    @staticmethod
    def extract_text_from_docx(file_path: str) -> str:
        """Extract text dari DOCX file."""
        try:
            logger.info(f"Extracting text from DOCX: {file_path}")

            doc = Document(file_path)
            text = ""

            for para in doc.paragraphs:
                if para.text.strip():
                    text += para.text + "\n"

            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        if cell.text.strip():
                            text += cell.text + " "
                    text += "\n"

            return text.strip()

        except Exception as e:
            logger.error(f"Error extracting text from DOCX: {str(e)}")
            raise

    @staticmethod
    def extract_text(file_path: str, file_type: str) -> str:
        """Extract text dari file berdasarkan tipe ('pdf' atau 'docx')."""
        if file_type.lower() == "pdf":
            return OCRService.extract_text_from_pdf(file_path)
        elif file_type.lower() == "docx":
            return OCRService.extract_text_from_docx(file_path)
        else:
            raise ValueError(f"Unsupported file type: {file_type}")


class CVParsingService:
    """Service untuk parsing CV data dari extracted text."""

    @staticmethod
    def extract_name(text: str) -> str:
        """Extract nama dari CV text."""
        lines = text.split("\n")[:5]

        for line in lines:
            line = line.strip()
            if len(line) > 3 and len(line) < 100 and line.isupper():
                return line

        for line in text.split("\n"):
            if line.strip():
                return line.strip()[:100]

        return "Unknown"

    @staticmethod
    def extract_email(text: str) -> str:
        """Extract email dari CV text."""
        email_pattern = r"[a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,}"
        matches = re.findall(email_pattern, text)
        return matches[0] if matches else ""

    @staticmethod
    def extract_phone(text: str) -> str:
        """Extract nomor telepon dari CV text."""
        phone_patterns = [
            r"\+62\s?[\d\s\-]{9,}",
            r"62[\d\s\-]{9,}",
            r"\+\d{1,3}[\d\s\-]{9,}",
            r"[\d\s\-]{10,}",
        ]

        for pattern in phone_patterns:
            matches = re.findall(pattern, text)
            if matches:
                return matches[0].replace(" ", "").replace("-", "")

        return ""

    @staticmethod
    def extract_experience_years(text: str) -> int:
        """Extract pengalaman kerja (tahun) dari CV text."""
        patterns = [
            r"(\d+)\s+(?:years?|yrs?)\s+(?:of\s+)?(?:work|experience)",
            r"(?:work|experience).*?(\d+)\s+(?:years?|yrs?)",
            r"(\d+)\s+year",
        ]

        for pattern in patterns:
            matches = re.findall(pattern, text, re.IGNORECASE)
            if matches:
                try:
                    years = int(matches[0])
                    if 0 <= years <= 70:
                        return years
                except ValueError:
                    pass

        return 0

    @staticmethod
    def extract_education(text: str) -> str:
        """Extract informasi pendidikan dari CV text."""
        education_keywords = [
            "bachelor",
            "master",
            "phd",
            "diploma",
            "b.s.",
            "m.s.",
            "b.a.",
            "m.a.",
            "degree",
            "university",
            "college",
        ]

        lines = text.split("\n")
        for line in lines:
            line_lower = line.lower()
            if any(keyword in line_lower for keyword in education_keywords):
                return line.strip()[:255]

        return ""

    @staticmethod
    def extract_skills(text: str) -> List[str]:
        """Extract skills dari CV text."""
        common_skills = {
            "languages": [
                "python", "javascript", "java", "c++", "c#", "php",
                "ruby", "go", "rust", "typescript", "sql",
            ],
            "frameworks": [
                "django", "flask", "fastapi", "react", "vue", "angular",
                "spring", "asp.net", "express",
            ],
            "databases": [
                "mysql", "postgresql", "mongodb", "redis", "elasticsearch",
                "dynamodb", "oracle",
            ],
            "tools": [
                "git", "docker", "kubernetes", "jenkins", "ci/cd", "aws",
                "gcp", "azure", "linux",
            ],
            "data": [
                "pandas", "numpy", "scikit-learn", "tensorflow", "pytorch",
                "spark", "hadoop",
            ],
        }

        found_skills = []
        text_lower = text.lower()

        for category_skills in common_skills.values():
            for skill in category_skills:
                if re.search(r"\b" + skill + r"\b", text_lower):
                    if skill not in found_skills:
                        found_skills.append(skill)

        return found_skills[:20]

    @staticmethod
    def parse_cv(text: str) -> Dict:
        """Parse CV text dan extract semua informasi."""
        logger.info("Parsing CV data...")

        parsed_data = {
            "name": CVParsingService.extract_name(text),
            "email": CVParsingService.extract_email(text),
            "phone": CVParsingService.extract_phone(text),
            "experience_years": CVParsingService.extract_experience_years(text),
            "education": CVParsingService.extract_education(text),
            "skills": CVParsingService.extract_skills(text),
            "cv_text": text[:5000],
        }

        logger.info(f"CV parsing completed: {parsed_data['name']}")
        return parsed_data


class CVProcessingService:
    """Service untuk complete CV processing pipeline."""

    @staticmethod
    def process_cv_file(file_path: str, file_type: str) -> Dict:
        """
        Process CV file: extract text dan parse data.

        Args:
            file_path: Path ke CV file
            file_type: Tipe file ('pdf' atau 'docx')

        Returns:
            Dictionary dengan parsed CV data
        """
        logger.info(f"Processing CV file: {file_path}")

        text = OCRService.extract_text(file_path, file_type)
        logger.info(f"Extracted {len(text)} characters from CV")

        return CVParsingService.parse_cv(text)